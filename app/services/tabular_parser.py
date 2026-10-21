"""表格类文档解析（11 号需求 FR-03/FR-04）

覆盖 .xlsx/.xlsm(openpyxl) / .xls(xlrd 2.x) / .csv(pandas + 编码探测)。

核心设计 = 双形态切片前置在这里做：
    - 数据表型 sheet(平均单元格 <50 字): 每 ROWS_PER_BLOCK 行一个段落,
      **每个段落重复表头行**,前缀【文件名 · sheet名】,保证 embedding 命中后
      LLM 看到的行永远带列名。
    - 文档型 sheet(平均单元格 >=50 字): 逐行句子化(表头列名：值),
      拼成连续正文交给现有滑窗切块。
    - 混合形态: 前部连续长文本行按文档型处理,其后表体按数据表型处理。

熔断(FR-04): 单 sheet 行数 MAX_SHEET_ROWS=50_000 → ValueError;
             workbook sheet 数 MAX_SHEETS=50 → 只解析前 50 并记录 partial。

输出与 parser.ParsedDoc 约定一致: 段落之间 "\\n\\n" 分隔。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, List, Optional

from app.core.logger import get_logger
from app.services.parser import ParsedDoc

logger = get_logger(__name__)

# ---- 熔断线(FR-04) ----
MAX_SHEET_ROWS = 50_000
MAX_SHEETS = 50
# ---- 数据表型 ----
ROWS_PER_BLOCK = 20
# ---- 形态判断阈值(FR-03: <20 数据 / >=50 文档 / 20-50 默认数据) ----
DOC_CELL_AVG = 50


class TabularError(ValueError):
    """表格解析类错误(沿用 worker 对 ValueError → 解析失败 的处理)。"""


# ---------------------------------------------------------------- 形态判断

def _row_avg_len(cells: List[str]) -> float:
    vals = [c for c in cells if c]
    if not vals:
        return 0.0
    return sum(len(c) for c in vals) / len(vals)


def _classify(rows: List[List[str]]) -> str:
    """返回 'doc' / 'data' / 'mixed'(mixed = 前部有长文本导言)。"""
    nonempty = [r for r in rows if any(c.strip() for c in r)]
    if not nonempty:
        return "empty"
    # 找"导言区": 开头连续的长文本行(平均 >=DOC_CELL_AVG),上限 30 行
    preface = 0
    for r in nonempty[:30]:
        if _row_avg_len(r) >= DOC_CELL_AVG:
            preface += 1
        else:
            break
    if preface == len(nonempty):
        return "doc"
    avg_all = sum(_row_avg_len(r) for r in nonempty) / len(nonempty)
    if avg_all >= DOC_CELL_AVG:
        return "doc"
    if preface >= 2:
        return "mixed"
    return "data"


# ---------------------------------------------------------------- 段落生成

def _data_block(filename: str, sheet: str, header: List[str],
                body: List[List[str]]) -> Iterable[str]:
    """数据表型: 每 ROWS_PER_BLOCK 行一块,块内重复表头。"""
    head_line = " | ".join(h or f"列{i+1}" for i, h in enumerate(header))
    prefix = f"【{filename} · {sheet}】"
    for start in range(0, len(body), ROWS_PER_BLOCK):
        block = body[start:start + ROWS_PER_BLOCK]
        lines = [f"{prefix}表头: {head_line}"]
        lines.extend(" | ".join(row) for row in block)
        yield "\n".join(lines)


def _docify(filename: str, sheet: str, header: List[str],
            rows: List[List[str]]) -> Iterable[str]:
    """文档型: 逐行句子化 `列名：值`,拼成连续正文。"""
    prefix = f"【{filename} · {sheet}】"
    for row in rows:
        pairs = []
        for i, cell in enumerate(row):
            if not cell.strip():
                continue
            col = header[i].strip() if i < len(header) and header[i].strip() else ""
            pairs.append(f"{col}：{cell}" if col else cell)
        if pairs:
            yield f"{prefix}" + "，".join(pairs)


def sheet_to_paragraphs(filename: str, sheet: str,
                        rows: List[List[str]]) -> List[str]:
    """单 sheet → 段落列表(核心切块规则,单测打这里)。"""
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []
    kind = _classify(rows)
    if kind == "doc":
        # 全文档型: 首行当表头做句子化
        header = rows[0] if _row_avg_len(rows[0]) < DOC_CELL_AVG * 3 else []
        body = rows[1:] if header else rows
        return list(_docify(filename, sheet, header, body))
    # data / mixed: 首行为表头
    preface_rows, table_rows = [], rows
    if kind == "mixed":
        split = 0
        for r in rows[:30]:
            if _row_avg_len(r) >= DOC_CELL_AVG:
                split += 1
            else:
                break
        preface_rows, table_rows = rows[:split], rows[split:]
    paras: List[str] = []
    if preface_rows:
        # 导言区句子化(无表头概念)
        for r in preface_rows:
            text = " ".join(c for c in r if c.strip())
            if text:
                paras.append(f"【{filename} · {sheet}】{text}")
    header = table_rows[0] if len(table_rows) >= 2 else []
    body = table_rows[1:] if len(table_rows) >= 2 else table_rows
    if header or body:
        paras.extend(_data_block(filename, sheet, header or [], body))
    return paras


# ---------------------------------------------------------------- openpyxl(.xlsx)

def _xlsx_rows(ws) -> List[List[str]]:
    """读 sheet 全部行;展开合并单元格(左上值填充整个区域)。"""
    if ws.max_row and ws.max_row > MAX_SHEET_ROWS:
        raise TabularError(
            f"工作表「{ws.title}」行数 {ws.max_row} 超过上限 {MAX_SHEET_ROWS},请拆分后上传")
    grid: List[List[str]] = []
    for row in ws.iter_rows(values_only=True):
        grid.append([_cell_str(v) for v in row])
    # 合并单元格展开
    merged = list(getattr(ws, "merged_cells", []) or [])
    if merged and grid:
        import copy
        area_map = copy.deepcopy(merged)
        for rng in area_map:
            min_col, min_row = rng.min_col, rng.min_row
            try:
                top_left = grid[min_row - 1][min_col - 1]
            except IndexError:
                continue
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    try:
                        grid[r - 1][c - 1] = top_left
                    except IndexError:
                        pass
    return grid


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(path: Path) -> ParsedDoc:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(path), data_only=True, read_only=False)
    except Exception as exc:
        raise TabularError(f"Excel 文件打不开(可能损坏或加密): {exc}") from exc

    sheets = wb.worksheets[:MAX_SHEETS]
    partial_note = ""
    if len(wb.worksheets) > MAX_SHEETS:
        partial_note = (f"共 {len(wb.worksheets)} 个 sheet,仅解析前 {MAX_SHEETS} 个")
        logger.warning("xlsx sheet 超限 partial: %s", path.name)

    name = path.name
    paragraphs: List[str] = []
    for ws in sheets:
        paragraphs.extend(sheet_to_paragraphs(name, ws.title, _xlsx_rows(ws)))
    wb.close()
    return _finish(paragraphs, path, "xlsx", partial_note)


# ---------------------------------------------------------------- xlrd(.xls)

def parse_xls(path: Path) -> ParsedDoc:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise TabularError(f"老版 Excel(.xls) 打不开: {exc}") from exc

    name = path.name
    paragraphs: List[str] = []
    for sheet in book.sheets()[:MAX_SHEETS]:
        if sheet.nrows > MAX_SHEET_ROWS:
            raise TabularError(
                f"工作表「{sheet.name}」行数 {sheet.nrows} 超过上限 {MAX_SHEET_ROWS},请拆分后上传")
        grid = [[_cell_str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]
        paragraphs.extend(sheet_to_paragraphs(name, sheet.name, grid))
    partial = "" if book.nsheets <= MAX_SHEETS else \
        f"共 {book.nsheets} 个 sheet,仅解析前 {MAX_SHEETS} 个"
    return _finish(paragraphs, path, "xls", partial)


# ---------------------------------------------------------------- csv(pandas)

_CSV_MAX_BYTES = 50 * 1024 * 1024  # csv 单文件 50MB 熔断(文本表不会更大才合理)


def parse_csv(path: Path) -> ParsedDoc:
    import pandas as pd

    size = path.stat().st_size
    if size > _CSV_MAX_BYTES:
        raise TabularError(f"CSV 文件 {size/1024/1024:.0f}MB 超过 50MB,请拆分后上传")

    raw = path.read_bytes()
    text: Optional[str] = None
    encoding = ""
    for enc in ("utf-8-sig", "gbk"):  # WPS 导出 csv 常见 gbk
        try:
            text = raw.decode(enc)
            encoding = enc
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise TabularError("CSV 编码无法识别(尝试 utf-8/gbk 均失败),请另存为 utf-8 CSV")

    from io import StringIO
    try:
        df = pd.read_csv(StringIO(text), dtype=str, keep_default_na=False,
                         nrows=MAX_SHEET_ROWS + 1)
    except Exception as exc:
        raise TabularError(f"CSV 解析失败: {exc}") from exc
    partial = ""
    if len(df) > MAX_SHEET_ROWS:
        raise TabularError(
            f"CSV 行数 {len(df)} 超过上限 {MAX_SHEET_ROWS},请拆分后上传")

    header = [str(c) for c in df.columns]
    grid = [header] + [[_cell_str(v) for v in row] for row in df.itertuples(index=False)]
    paras = sheet_to_paragraphs(path.name, "CSV", grid)
    meta_extra = f"csv 编码 {encoding}"
    return _finish(paras, path, "csv", (partial + " " + meta_extra).strip())


# ---------------------------------------------------------------- 公共收尾

def _finish(paragraphs: List[str], path: Path, fmt: str,
            partial_note: str = "") -> ParsedDoc:
    paragraphs = [p for p in paragraphs if p.strip()]
    if not paragraphs:
        raise TabularError(f"未能从 {path.name} 提取到任何文本内容(空表或仅格式)")
    content = "\n\n".join(paragraphs)
    metadata = {
        "filename": path.name,
        "format": fmt,
        "parse_method": f"tabular:{fmt}",
        "blocks": len(paragraphs),
    }
    if partial_note:
        metadata["partial_note"] = partial_note
    logger.info("表格解析完成: %s fmt=%s blocks=%d chars=%d",
                path.name, fmt, len(paragraphs), len(content))
    return ParsedDoc(content=content, metadata=metadata)


def parse_tabular(path: Path) -> ParsedDoc:
    """表格路由统一入口(按后缀分发)。"""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return parse_xlsx(path)
    if suffix == ".xls":
        return parse_xls(path)
    if suffix == ".csv":
        return parse_csv(path)
    raise TabularError(f"不是表格文件: {suffix}")
