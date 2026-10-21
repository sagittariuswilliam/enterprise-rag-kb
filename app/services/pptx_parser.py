"""PPTX 解析(11 号需求 FR-01): python-pptx 逐页提取文本框 + 表格。

输出约定与 parser 一致: 每页一段,段落间 "\\n\\n";页前缀【页 N】,
让 chunker 的按页/按段逻辑自然可用。表格走 "| " 连接(同 docx 表格标记风格)。
熔断: >MAX_PPTX_SLIDES 页拒收。
"""

from __future__ import annotations

from pathlib import Path

from app.core.logger import get_logger
from app.services.parser import ParsedDoc

logger = get_logger(__name__)

MAX_PPTX_SLIDES = 200


class PptxError(ValueError):
    pass


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cells.append(cell.text.strip().replace("\n", " "))
        if any(cells):
            rows.append(" | ".join(cells))
    if not rows:
        return ""
    return "[表格]\n" + "\n".join(rows)


def parse_pptx(path: Path) -> ParsedDoc:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise PptxError(f"PPT 文件打不开(损坏或老 .ppt 格式): {exc}") from exc

    slides = list(prs.slides)
    if len(slides) > MAX_PPTX_SLIDES:
        raise PptxError(
            f"PPT 共 {len(slides)} 页,超过上限 {MAX_PPTX_SLIDES},请拆分后上传")

    sections = []
    for idx, slide in enumerate(slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    parts.append(txt)
            if getattr(shape, "has_table", False):
                tb = _table_text(shape.table)
                if tb:
                    parts.append(tb)
        # 备注页也提(常藏讲解文字)
        try:
            if slide.has_notes_slide:
                note = slide.notes_slide.notes_text_frame.text.strip()
                if note:
                    parts.append(f"[备注] {note}")
        except Exception:
            pass
        if parts:
            sections.append(f"【页 {idx}】\n" + "\n".join(parts))

    if not sections:
        raise PptxError(f"未能从 {path.name} 提取到任何文本(纯图片 PPT 请先导出为 PDF)")
    content = "\n\n".join(sections)
    return ParsedDoc(content, {
        "filename": path.name,
        "format": "pptx",
        "parse_method": "python-pptx",
        "pages": len(slides),
    })
